import os
import pandas as pd
import numpy as np
import nbformat as nbf
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from scipy import stats

print("==================================================================")
print("     COMPREHENSIVE FINAL QUALITY-CONTROL & VERIFICATION PASS      ")
print("==================================================================")

# 1. DATA AUDIT
print("\n[CHECK 1] Dataset Integrity Audit:")
raw_csv = 'seasonal_agriculture_performance_dataset.csv'
clean_csv = 'seasonal_agriculture_performance_cleaned.csv'

assert os.path.exists(raw_csv), "Raw CSV missing!"
assert os.path.exists(clean_csv), "Cleaned CSV missing!"

raw_df = pd.read_csv(raw_csv)
clean_df = pd.read_csv(clean_csv)

print(f"  Raw Dataset Dimensions:    {raw_df.shape[0]} rows, {raw_df.shape[1]} columns")
assert raw_df.shape == (4000, 28), f"Unexpected raw shape: {raw_df.shape}"

print(f"  Cleaned Dataset Dimensions:{clean_df.shape[0]} rows, {clean_df.shape[1]} columns")
assert clean_df.shape == (4000, 29), f"Unexpected cleaned shape: {clean_df.shape}"

raw_missing = raw_df.isnull().sum()
print(f"  Raw Missing Values:        Rainfall={raw_missing['Rainfall_mm']}, Soil_Moisture={raw_missing['Soil_Moisture_pct']}, Yield={raw_missing['Yield_Tonnes_Ha']}")
assert raw_missing['Rainfall_mm'] == 48
assert raw_missing['Soil_Moisture_pct'] == 40
assert raw_missing['Yield_Tonnes_Ha'] == 32

clean_missing = clean_df.isnull().sum().sum()
print(f"  Cleaned Missing Values:    {clean_missing} (Must be exactly 0)")
assert clean_missing == 0, f"Remaining nulls in clean_df: {clean_missing}"

dup_rows = clean_df.duplicated().sum()
dup_ids = clean_df['Farm_ID'].duplicated().sum()
print(f"  Duplicate Rows:            {dup_rows}")
print(f"  Duplicate Farm_IDs:        {dup_ids}")
assert dup_rows == 0, f"Found {dup_rows} duplicate rows!"
assert dup_ids == 0, f"Found {dup_ids} duplicate Farm_IDs!"

# 2. EXACT FIVE CHARTS
print("\n[CHECK 2] Exact Five Main Result Charts:")
expected_charts = [
    'chart1_seasonal_yield.png',
    'chart2_seasonal_profit.png',
    'chart3_seasonal_rainfall.png',
    'chart4_crop_season_heatmap.png',
    'chart5_rainfall_yield_scatter.png'
]
for c in expected_charts:
    assert os.path.exists(c), f"Chart missing: {c}"
    size = os.path.getsize(c)
    print(f"  PASS: {c:<36} ({size:>7,} bytes)")

# 3. STATISTICAL CALCULATIONS
print("\n[CHECK 3] Statistical Rigor & Non-Parametric Metrics:")
# Kruskal-Wallis
ky = clean_df[clean_df['Season'] == 'Kharif']['Yield_Tonnes_Ha']
ry = clean_df[clean_df['Season'] == 'Rabi']['Yield_Tonnes_Ha']
zy = clean_df[clean_df['Season'] == 'Zaid']['Yield_Tonnes_Ha']
kw_stat, kw_p = stats.kruskal(ky, ry, zy)
print(f"  Kruskal-Wallis H-statistic: {kw_stat:.4f}")
print(f"  Kruskal-Wallis p-value:     {kw_p:.4e}")
assert abs(kw_stat - 70.5935) < 0.01

# Spearman Correlation
rho, p_val = stats.spearmanr(clean_df['Rainfall_mm'], clean_df['Yield_Tonnes_Ha'])
print(f"  Spearman Rho (rho):         {rho:.4f}")
print(f"  Spearman p-value:           {p_val:.4e}")
print(f"  Sample Size (N):            {len(clean_df)}")
assert abs(rho - 0.1295) < 0.01

# 4. NOTEBOOK AUDIT
print("\n[CHECK 4] Jupyter Notebook Audit:")
nb_path = 'Seasonal_Agriculture_Performance_Analysis.ipynb'
assert os.path.exists(nb_path), f"Notebook missing: {nb_path}"
with open(nb_path, 'r', encoding='utf-8') as f:
    nb = nbf.read(f, as_version=4)

code_cells = [c for c in nb.cells if c.cell_type == 'code']
executed = [c for c in code_cells if len(c.outputs) > 0]
print(f"  Total Cells in Notebook:    {len(nb.cells)}")
print(f"  Code Cells Executed:        {len(executed)} / {len(code_cells)}")
assert len(executed) == len(code_cells), "Unexecuted code cells found!"

# Check no cell has error outputs
for idx, c in enumerate(code_cells):
    has_err = any(out.output_type == 'error' for out in c.outputs)
    assert not has_err, f"Error in cell index {idx}!"
print("  All code cells executed without errors.")

# 5. POWERPOINT AUDIT
print("\n[CHECK 5] PowerPoint Submission Deck Audit:")
ppt_path = 'VOIS_Major_Project_PPT_Submission_Template.pptx'
assert os.path.exists(ppt_path), f"PPT missing: {ppt_path}"
prs = pptx.Presentation(ppt_path)
print(f"  Total Slides:               {len(prs.slides)} (Must be 14)")
assert len(prs.slides) == 14, f"Slide count is {len(prs.slides)}, expected 14!"

expected_titles = {
    5: "Seasonal Yield Performance",
    6: "Seasonal Economic Performance",
    7: "Seasonal Environmental Conditions",
    8: "Crop Performance Across Seasons",
    9: "Rainfall and Yield Relationship"
}

for s_idx, exp_title in expected_titles.items():
    slide = prs.slides[s_idx]
    slide_text = " ".join([s.text for s in slide.shapes if s.has_text_frame])
    assert exp_title in slide_text, f"Slide {s_idx+1} missing title: {exp_title}"
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.top > pptx.util.Inches(1.0) and s.top < pptx.util.Inches(6.8)]
    print(f"  Slide {s_idx+1:2d} ({exp_title}): Title verified, {len(pics)} result chart embedded.")
    assert len(pics) >= 1, f"Slide {s_idx+1} missing chart image!"

# Check restricted words
print("\n[CHECK 6] Academic Terminology & Non-Causality Check:")
restricted_words = ['biomass', 'python 3.14']
with open('README.md', 'r', encoding='utf-8') as f:
    rm_text = f.read().lower()
for w in restricted_words:
    assert w not in rm_text, f"Found '{w}' in README.md!"

all_ppt_text = ""
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            all_ppt_text += " " + sh.text.lower()
for w in restricted_words:
    assert w not in all_ppt_text, f"Found '{w}' in PPT!"

print("  No prohibited 'biomass' or 'Python 3.14' claims found in README or PPT.")
print("\n==================================================================")
print("     ALL QUALITY-CONTROL AUDIT CHECKS PASSED SUCCESSFULLY (100%)    ")
print("==================================================================")
